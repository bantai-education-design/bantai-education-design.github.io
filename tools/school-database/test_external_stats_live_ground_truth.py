#!/usr/bin/env python3
"""Strict live Ground Truth audit for two external education indicators.

Sources checked on every run:
- MEXT / e-Stat, FY2024 School ICT survey, Table 2 (1) Total
  statInfId=000040365967
- Children and Families Agency, daycare-related status summary as of 2025-04-01
  official workbook "資料1～6"

The parser identifies the semantic metric column from official workbook headers and
then compares all 47 prefectures against the committed external-stat master JSON.
There is no local-file fallback in this audit.
"""

from __future__ import annotations

import io
import json
import re
import urllib.request
from pathlib import Path
from typing import Iterable

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
CARD_METADATA_PATH = ROOT / "data" / "school-database" / "prefecture-card-metadata.json"
EXT_STATS_PATH = ROOT / "data" / "school-database" / "prefecture-education-external-stats.json"

ICT_STAT_ID = "000040365967"
ICT_URL = f"https://www.e-stat.go.jp/stat-search/file-download?fileKind=0&statInfId={ICT_STAT_ID}"
WAITING_CHILDREN_URL = (
    "https://www.cfa.go.jp/assets/contents/node/basic_page/field_ref_resources/"
    "b0a8057b-34bf-4c20-84fb-ae592708ca9b/1a728dcc/"
    "20250828_policies_hoiku_torimatome_r7_02.xlsx"
)

PREFECTURE_NAMES = [
    "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県",
    "茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県",
    "新潟県", "富山県", "石川県", "福井県", "山梨県", "長野県", "岐阜県",
    "静岡県", "愛知県", "三重県", "滋賀県", "京都府", "大阪府", "兵庫県",
    "奈良県", "和歌山県", "鳥取県", "島根県", "岡山県", "広島県", "山口県",
    "徳島県", "香川県", "愛媛県", "高知県", "福岡県", "佐賀県", "長崎県",
    "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県",
]
PREFECTURE_SET = set(PREFECTURE_NAMES)


def _norm(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u3000", " ").strip()
    return re.sub(r"\s+", "", text)


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _norm(value).replace(",", "").replace("％", "%").replace("%", "")
    text = text.replace("▲", "-").replace("△", "-")
    if not text or text in {"-", "－", "…", "―"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _fetch(url: str, label: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as response:
        data = response.read()
    assert data[:2] == b"PK", f"{label}: downloaded payload is not an XLSX zip"
    print(f"[Official Live Audit] {label} downloaded live ({len(data):,} bytes)")
    return data


def _merged_header_text(sheet, col: int, first_data_row: int) -> str:
    parts: list[str] = []
    start = max(1, first_data_row - 12)
    end = first_data_row - 1
    for row in range(start, end + 1):
        value = sheet.cell(row, col).value
        if value is not None:
            parts.append(str(value))
    for merged in sheet.merged_cells.ranges:
        if merged.min_col <= col <= merged.max_col and merged.min_row <= end and merged.max_row >= start:
            value = sheet.cell(merged.min_row, merged.min_col).value
            if value is not None:
                parts.append(str(value))
    return _norm(" ".join(parts))


def _rows_by_prefecture(sheet) -> dict[str, list[int]]:
    rows: dict[str, list[int]] = {name: [] for name in PREFECTURE_NAMES}
    max_scan_col = min(sheet.max_column, 20)
    for r in range(1, sheet.max_row + 1):
        for c in range(1, max_scan_col + 1):
            text = _norm(sheet.cell(r, c).value)
            if text in PREFECTURE_SET:
                rows[text].append(r)
                break
    return rows


def _extract_prefecture_column(
    workbook,
    *,
    header_tokens: Iterable[str],
    expected_national: float | None = None,
    expected_sum: float | None = None,
) -> tuple[dict[str, float], str, int]:
    """Find one semantic metric column containing all 47 prefectures.

    Candidate columns must have a nearby header containing every requested token.
    If expected_national is supplied, a nearby 全国 row in the same column must match.
    If expected_sum is supplied, the 47 extracted values must sum to that value.
    """
    tokens = tuple(_norm(t) for t in header_tokens)
    diagnostics: list[str] = []

    for sheet in workbook.worksheets:
        pref_rows = _rows_by_prefecture(sheet)
        if any(not rows for rows in pref_rows.values()):
            continue

        # Workbooks may repeat prefecture tables. Try each occurrence index as one table block.
        max_occurrences = max(len(rows) for rows in pref_rows.values())
        for occ in range(max_occurrences):
            selected_rows: dict[str, int] = {}
            for name, rows in pref_rows.items():
                if occ < len(rows):
                    selected_rows[name] = rows[occ]
            if len(selected_rows) != 47:
                continue

            first_data_row = min(selected_rows.values())
            last_data_row = max(selected_rows.values())
            if last_data_row - first_data_row > 80:
                continue

            for col in range(1, sheet.max_column + 1):
                header = _merged_header_text(sheet, col, first_data_row)
                if not all(token in header for token in tokens):
                    continue

                values: dict[str, float] = {}
                for name, row in selected_rows.items():
                    value = _number(sheet.cell(row, col).value)
                    if value is None:
                        break
                    values[name] = value
                if len(values) != 47:
                    continue

                if expected_sum is not None and abs(sum(values.values()) - expected_sum) > 1e-9:
                    diagnostics.append(
                        f"{sheet.title}!col{col}: semantic header matched but sum={sum(values.values())}"
                    )
                    continue

                if expected_national is not None:
                    national_match = False
                    for r in range(max(1, first_data_row - 5), min(sheet.max_row, last_data_row + 5) + 1):
                        row_has_national = any(
                            _norm(sheet.cell(r, c).value) in {"全国", "全国計", "計全国"}
                            for c in range(1, min(sheet.max_column, 20) + 1)
                        )
                        if not row_has_national:
                            continue
                        national_value = _number(sheet.cell(r, col).value)
                        if national_value is not None and abs(national_value - expected_national) < 1e-9:
                            national_match = True
                            break
                    if not national_match:
                        diagnostics.append(
                            f"{sheet.title}!col{col}: semantic header matched but national {expected_national} not found"
                        )
                        continue

                return values, sheet.title, col

    detail = "; ".join(diagnostics[-8:]) if diagnostics else "no semantic candidate column found"
    raise AssertionError(f"Unable to identify official metric column: {detail}")


def build_live_official_external_stats() -> dict[str, dict[str, float]]:
    ict_bytes = _fetch(ICT_URL, f"MEXT/e-Stat ICT Table 2 (1), statInfId={ICT_STAT_ID}")
    ict_wb = openpyxl.load_workbook(io.BytesIO(ict_bytes), data_only=True)
    ict_by_name, ict_sheet, ict_col = _extract_prefecture_column(
        ict_wb,
        header_tokens=("教材研究", "ICT"),
        expected_national=90.7,
    )
    assert all(0 <= value <= 100 for value in ict_by_name.values())
    print(f"[Official Live Audit] ICT Category A identified at {ict_sheet}!col{ict_col}")

    waiting_bytes = _fetch(WAITING_CHILDREN_URL, "CFA waiting-children workbook (2025-04-01)")
    waiting_wb = openpyxl.load_workbook(io.BytesIO(waiting_bytes), data_only=True)
    waiting_by_name, waiting_sheet, waiting_col = _extract_prefecture_column(
        waiting_wb,
        header_tokens=("待機児童",),
        expected_sum=2254.0,
    )
    assert all(value >= 0 and float(value).is_integer() for value in waiting_by_name.values())
    print(f"[Official Live Audit] waiting children identified at {waiting_sheet}!col{waiting_col}")

    return {
        name: {
            "ict_teaching_capability": round(ict_by_name[name], 1),
            "waiting_children_count": int(waiting_by_name[name]),
        }
        for name in PREFECTURE_NAMES
    }


def test_external_stats_live_ground_truth() -> None:
    card_payload = json.loads(CARD_METADATA_PATH.read_text(encoding="utf-8"))
    ext_payload = json.loads(EXT_STATS_PATH.read_text(encoding="utf-8"))

    name_to_slug = {
        entry["prefecture_name"]: entry["prefecture_code"]
        for entry in card_payload["prefectures"]
    }
    assert set(name_to_slug) == PREFECTURE_SET
    ext_by_slug = {entry["prefecture_code"]: entry for entry in ext_payload["prefectures"]}
    assert len(ext_by_slug) == 47

    live = build_live_official_external_stats()
    assert len(live) == 47
    assert sum(item["waiting_children_count"] for item in live.values()) == 2254

    for name, live_entry in live.items():
        slug = name_to_slug[name]
        committed = ext_by_slug[slug]
        assert live_entry["ict_teaching_capability"] == committed["ict_teaching_capability"], (
            f"{name}: ICT live={live_entry['ict_teaching_capability']} committed={committed['ict_teaching_capability']}"
        )
        assert live_entry["waiting_children_count"] == committed["waiting_children_count"], (
            f"{name}: waiting live={live_entry['waiting_children_count']} committed={committed['waiting_children_count']}"
        )


if __name__ == "__main__":
    test_external_stats_live_ground_truth()
    print("ALL 47 PREFECTURES ICT & WAITING-CHILDREN LIVE OFFICIAL GROUND TRUTH TESTS PASSED SUCCESSFULLY!")
