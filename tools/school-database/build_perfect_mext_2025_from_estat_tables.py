import urllib.request
import openpyxl
from io import BytesIO
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUTPUT_PATH = ROOT / 'data' / 'school-database' / 'mext-school-basic-survey-2025-official.json'

PREFECTURE_MAP = [
    ('hokkaido', '北海道'), ('aomori', '青森県'), ('iwate', '岩手県'), ('miyagi', '宮城県'),
    ('akita', '秋田県'), ('yamagata', '山形県'), ('fukushima', '福島県'), ('ibaraki', '茨城県'),
    ('tochigi', '栃木県'), ('gunma', '群馬県'), ('saitama', '埼玉県'), ('chiba', '千葉県'),
    ('tokyo', '東京都'), ('kanagawa', '神奈川県'), ('niigata', '新潟県'), ('toyama', '富山県'),
    ('ishikawa', '石川県'), ('fukui', '福井県'), ('yamanashi', '山梨県'), ('nagano', '長野県'),
    ('gifu', '岐阜県'), ('shizuoka', '静岡県'), ('aichi', '愛知県'), ('mie', '三重県'),
    ('shiga', '滋賀県'), ('kyoto', '京都府'), ('osaka', '大阪府'), ('hyogo', '兵庫県'),
    ('nara', '奈良県'), ('wakayama', '和歌山県'), ('tottori', '鳥取県'), ('shimane', '島根県'),
    ('okayama', '岡山県'), ('hiroshima', '広島県'), ('yamaguchi', '山口県'), ('tokushima', '徳島県'),
    ('kagawa', '香川県'), ('ehime', '愛媛県'), ('kochi', '高知県'), ('fukuoka', '福岡県'),
    ('saga', '佐賀県'), ('nagasaki', '長崎県'), ('kumamoto', '熊本県'), ('oita', '大分県'),
    ('miyazaki', '宮崎県'), ('kagoshima', '鹿児島県'), ('okinawa', '沖縄県')
]

NAME_TO_CODE = {name: code for code, name in PREFECTURE_MAP}

TABLE_DOWNLOAD_IDS = {
    'elem_schools': '000040393694',  # 表41: 都道府県別学校数
    'elem_classes': '000040393699',  # 表46: 都道府県別編制方式別学級数
    'elem_students': '000040393701', # 表48: 都道府県別学年別児童数
    'elem_teachers': '000040393707', # 表54: 都道府県別職名別教員数（本務者）
    'jhs_schools': '000040393720',   # 表67: 都道府県別学校数
    'jhs_classes': '000040393726',   # 表73: 都道府県別編制方式別学級数
    'jhs_students': '000040393728',  # 表75: 都道府県別学年別生徒数
    'jhs_teachers': '000040393735',  # 表82: 都道府県別職名別教員数（本務者）
    'special_needs_schools': '000040393852' # 表199: 特別支援学校・都道府県別学校数
}

def parse_table_data_exact(stat_id):
    url = f"https://www.e-stat.go.jp/stat-search/file-download?statInfId={stat_id}&fileKind=0"
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req) as resp:
        content = resp.read()
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    sheet = wb.active
    
    rows_data = {}
    for r in range(1, sheet.max_row + 1):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
        row_str = " ".join([str(v) for v in row_vals if v is not None])
        
        # Must be R7 / 令和7年度 row
        if not ('令和7' in row_str or 'R7' in row_str or '7' in str(row_vals[0])):
            continue
            
        # Match prefecture by exact cell equality
        matched_code = None
        for cell in row_vals[:5]: # Prefecture cell is in the first 5 columns
            if cell is None:
                continue
            cell_clean = str(cell).strip()
            if cell_clean in NAME_TO_CODE:
                matched_code = NAME_TO_CODE[cell_clean]
                break
                
        if matched_code:
            assert matched_code not in rows_data, f"Duplicate prefecture row detected for {matched_code} in table {stat_id}"
            rows_data[matched_code] = row_vals

    # Strict Validation Assertions
    assert len(rows_data) == 47, f"Expected 47 prefectures for table {stat_id}, but found {len(rows_data)}"
    assert set(rows_data.keys()) == set(code for code, name in PREFECTURE_MAP), f"Missing prefecture codes in table {stat_id}"
    
    return rows_data

def build_official_dataset_dict():
    parsed_tables = {}
    for key, sid in TABLE_DOWNLOAD_IDS.items():
        parsed_tables[key] = parse_table_data_exact(sid)
        
    pref_result_list = []
    for code, name in PREFECTURE_MAP:
        # 1. Elem schools (表41)
        row_es = parsed_tables['elem_schools'][code]
        elem_schools = int(row_es[2]) if len(row_es) > 2 and row_es[2] is not None else 0
        private_elem_schools = int(row_es[24]) if len(row_es) > 24 and row_es[24] is not None else 0
        
        # 2. Elem classes (表46)
        row_ec = parsed_tables['elem_classes'][code]
        elem_classes = int(row_ec[3]) if len(row_ec) > 3 and row_ec[3] is not None else 0
        
        # 3. Elem students (表48)
        row_est = parsed_tables['elem_students'][code]
        elem_students = int(row_est[2]) if len(row_est) > 2 and row_est[2] is not None else 0
        
        # 4. Elem teachers (表54)
        row_et = parsed_tables['elem_teachers'][code]
        elem_teachers = int(row_et[2]) if len(row_et) > 2 and row_et[2] is not None else 0
        
        # 5. JHS schools (表67)
        row_js = parsed_tables['jhs_schools'][code]
        jhs_schools = int(row_js[2]) if len(row_js) > 2 and row_js[2] is not None else 0
        
        # 6. JHS classes (表73)
        row_jc = parsed_tables['jhs_classes'][code]
        jhs_classes = int(row_jc[3]) if len(row_jc) > 3 and row_jc[3] is not None else 0
        
        # 7. JHS students (表75)
        row_jst = parsed_tables['jhs_students'][code]
        jhs_students = int(row_jst[2]) if len(row_jst) > 2 and row_jst[2] is not None else 0
        
        # 8. JHS teachers (表82)
        row_jt = parsed_tables['jhs_teachers'][code]
        jhs_teachers = int(row_jt[2]) if len(row_jt) > 2 and row_jt[2] is not None else 0
        
        # 9. Special needs schools (表199)
        row_sn = parsed_tables['special_needs_schools'][code]
        special_needs_schools = int(row_sn[2]) if len(row_sn) > 2 and row_sn[2] is not None else 0
        
        # Sanity check essential non-zero counts
        assert elem_schools > 0, f"Zero elem_schools for {name}"
        assert elem_students > 0, f"Zero elem_students for {name}"
        assert elem_teachers > 0, f"Zero elem_teachers for {name}"
        assert elem_classes > 0, f"Zero elem_classes for {name}"
        assert jhs_schools > 0, f"Zero jhs_schools for {name}"
        assert jhs_students > 0, f"Zero jhs_students for {name}"
        assert jhs_teachers > 0, f"Zero jhs_teachers for {name}"
        assert jhs_classes > 0, f"Zero jhs_classes for {name}"
        assert special_needs_schools > 0, f"Zero special_needs_schools for {name}"
        
        entry = {
            'prefecture_code': code,
            'prefecture_name': name,
            'elem_schools': elem_schools,
            'private_elem_schools': private_elem_schools,
            'elem_students': elem_students,
            'elem_teachers': elem_teachers,
            'elem_classes': elem_classes,
            'jhs_schools': jhs_schools,
            'jhs_students': jhs_students,
            'jhs_teachers': jhs_teachers,
            'jhs_classes': jhs_classes,
            'special_needs_schools': special_needs_schools
        }
        pref_result_list.append(entry)

    official_dataset = {
        'schema_version': '1.0',
        'title': '文部科学省 令和7年度学校基本調査 確定値（2025年12月26日公表・47都道府県完全転記正本）',
        'reference_date': '2025年5月1日現在',
        'published_date': '2025年12月26日',
        'citation': {
            'elem_schools': '文部科学省「令和7年度学校基本調査 確定値」表41（都道府県別学校数）',
            'elem_classes': '文部科学省「令和7年度学校基本調査 確定値」表46（都道府県別編制方式別学級数）',
            'elem_students': '文部科学省「令和7年度学校基本調査 確定値」表48（都道府県別学年別児童数）',
            'elem_teachers': '文部科学省「令和7年度学校基本調査 確定値」表54（都道府県別職名別教員数・本務者）',
            'jhs_schools': '文部科学省「令和7年度学校基本調査 確定値」表67（都道府県別学校数）',
            'jhs_classes': '文部科学省「令和7年度学校基本調査 確定値」表73（都道府県別編制方式別学級数）',
            'jhs_students': '文部科学省「令和7年度学校基本調査 確定値」表75（都道府県別学年別生徒数）',
            'jhs_teachers': '文部科学省「令和7年度学校基本調査 確定値」表82（都道府県別職名別教員数・本務者）',
            'special_needs_schools': '文部科学省「令和7年度学校基本調査 確定値」表199（特別支援学校・都道府県別学校数）'
        },
        'prefectures': pref_result_list
    }
    return official_dataset

def main():
    print("Downloading and parsing all 9 official e-Stat tables with exact cell matching...", flush=True)
    dataset = build_official_dataset_dict()
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(dataset, f, ensure_ascii=False, indent=2)
    print(f"SUCCESS: Generated {OUTPUT_PATH} with 47 prefectures cleanly parsed!", flush=True)

if __name__ == '__main__':
    main()
