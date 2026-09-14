#!/usr/bin/env python3
import io
import unicodedata
import urllib.request
import openpyxl

URL = "https://www.ipss.go.jp/pp-shicyoson/j/shicyoson23/2gaiyo_hyo/kekkahyo2_1.xlsx"
req = urllib.request.Request(URL, headers={"User-Agent":"Mozilla/5.0"})
with urllib.request.urlopen(req, timeout=60) as response:
    data=response.read()
wb=openpyxl.load_workbook(io.BytesIO(data), data_only=True)
print('sheets=', wb.sheetnames)
for ws in wb.worksheets:
    print(f'SHEET {ws.title!r} size={ws.max_row}x{ws.max_column}')
    shown=0
    for r in range(1, min(ws.max_row, 180)+1):
        vals=[]
        for c in range(1, ws.max_column+1):
            v=ws.cell(r,c).value
            if v is None or str(v).strip()=="":
                continue
            text=unicodedata.normalize('NFKC',str(v)).replace('\n',' ')
            vals.append(f'c{c}={text!r}')
        rowtext=' '.join(vals)
        if vals and (r <= 30 or '北海道' in rowtext or '東京都' in rowtext or '沖縄県' in rowtext or '0~14' in rowtext or '0～14' in rowtext or '指数' in rowtext or '2035' in rowtext or '2050' in rowtext):
            print(f'row{r}: ' + ' | '.join(vals))
            shown += 1
        if shown >= 90:
            break
