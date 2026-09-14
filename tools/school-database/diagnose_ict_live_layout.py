#!/usr/bin/env python3
import io
import unicodedata
import urllib.request
import openpyxl

ICT_URL = "https://www.e-stat.go.jp/stat-search/file-download?fileKind=0&statInfId=000040365967"
WAITING_URL = (
    "https://www.cfa.go.jp/assets/contents/node/basic_page/field_ref_resources/"
    "b0a8057b-34bf-4c20-84fb-ae592708ca9b/1a728dcc/"
    "20250828_policies_hoiku_torimatome_r7_02.xlsx"
)


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as response:
        return response.read()


def fmt(v):
    return unicodedata.normalize("NFKC", str(v)).replace("\n", " ")


ict = openpyxl.load_workbook(io.BytesIO(fetch(ICT_URL)), data_only=True)
print("ICT workbook sheets:", ict.sheetnames)
for ws in ict.worksheets:
    print(f"ICT SHEET {ws.title!r} size={ws.max_row}x{ws.max_column}")
    for r in list(range(1, min(ws.max_row, 9) + 1)) + list(range(max(1, ws.max_row - 4), ws.max_row + 1)):
        vals = [f"c{c}={fmt(ws.cell(r,c).value)!r}" for c in range(1, ws.max_column + 1) if ws.cell(r,c).value is not None and str(ws.cell(r,c).value).strip()]
        if vals:
            print(f"ICT row{r}: " + " | ".join(vals))

waiting = openpyxl.load_workbook(io.BytesIO(fetch(WAITING_URL)), data_only=True)
print("WAITING workbook sheets:", waiting.sheetnames)
for ws in waiting.worksheets:
    print(f"WAITING SHEET {ws.title!r} size={ws.max_row}x{ws.max_column}")
    shown = 0
    for r in range(1, min(ws.max_row, 40) + 1):
        vals = [f"c{c}={fmt(ws.cell(r,c).value)!r}" for c in range(1, min(ws.max_column, 30) + 1) if ws.cell(r,c).value is not None and str(ws.cell(r,c).value).strip()]
        rowtext = " ".join(vals)
        if vals and (r <= 12 or "北海道" in rowtext or "全国" in rowtext or "待機児童" in rowtext):
            print(f"WAITING row{r}: " + " | ".join(vals))
            shown += 1
        if shown >= 20:
            break
