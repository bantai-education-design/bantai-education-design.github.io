#!/usr/bin/env python3
"""Live official Ground Truth audit for IPSS 0-14 population indices.

Every CI run downloads the official IPSS Excel result table directly and compares
2035/2050 indices (2020=100) for all 47 prefectures with the committed canonical
projection JSON. There is no local-source fallback.
"""

from __future__ import annotations

import io
import json
import re
import unicodedata
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
MASTER_PATH = ROOT / "data" / "school-database" / "ipss-child-population-projection-2023.json"

IPSS_XLSX_URL = (
    "https://www.ipss.go.jp/pp-shicyoson/j/shicyoson23/"
    "2gaiyo_hyo/kekkahyo2_1.xlsx"
)

PREFECTURE_NAMES = [
    "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県",
    "茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県",
    "新潟県", "富山県", "石川県", "福井県", "山梨県", "長野県", "岐阜県",
    "静岡県", "愛知県", "三重県", "滋賀県", "京都府", "大阪府", "兵庫県",
    "奈良県", "和歌山県", "鳥取県", "島根県", "岡山県", "広島県", "山口県",
    "徳島県", "香川県", "愛媛県", "高知県", "福岡県", "佐賀県", "長崎県",
    "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県",
]
PREFECTURE_SET = set(PREFECTURE_NAMES)


def _norm(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\u3000", " ").replace("\n", " ").strip()
    return re.sub(r"\s+", "", text)


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _norm(value).replace(",", "").replace("%", "")
    if not text or text in {"-", "－", "…", "―"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _fetch_workbook() -> openpyxl.Workbook:
    request = urllib.request.Request(IPSS_XLSX_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=60) as response:
        data = response.read()
    assert data[:2] == b"PK", "IPSS live source is not an XLSX payload"
    print(f"[IPSS Live Audit] official result table downloaded live ({len(data):,} bytes)")
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)


def _header_text(sheet, col: int, first_data_row: int) -> str:
    start = max(1, first_data_row - 20)
    end = first_data_row - 1
    parts: list[str] = []
    for row in range(start, end + 1):
        value = sheet.cell(row, col).value
        if value is not None:
            parts.append(str(value))
    for merged in sheet.merged_cells.ranges:
        if merged.min_col <= col <= merged.max_col and merged.min_row <= end and merged.max_row >= start:
            value = sheet.cell(merged.min_row, merged.min_col).value
            if value is not None:
                parts.append(str(value))
    return _norm(" ".join(parts))


def _find_prefecture_block(workbook) -> tuple[object, int, dict[str, int]]:
    diagnostics: list[str] = []
    for sheet in workbook.worksheets:
        for col in range(1, sheet.max_column + 1):
            rows: dict[str, int] = {}
            duplicate = False
            for row in range(1, sheet.max_row + 1):
                name = _norm(sheet.cell(row, col).value)
                if name not in PREFECTURE_SET:
                    continue
                if name in rows:
                    duplicate = True
                    break
                rows[name] = row
            diagnostics.append(f"{sheet.title}!col{col}: prefectures={len(rows)} duplicate={duplicate}")
            if not duplicate and set(rows) == PREFECTURE_SET:
                row_numbers = sorted(rows.values())
                # The 47 prefectures in the official result table form one compact block.
                if row_numbers[-1] - row_numbers[0] <= 60:
                    return sheet, col, rows
    raise AssertionError("Unable to identify the official 47-prefecture block: " + "; ".join(diagnostics[-12:]))


def _find_index_column(sheet, first_data_row: int, year: int) -> int:
    candidates: list[int] = []
    year_tokens = {str(year), "令和17" if year == 2035 else "令和32"}
    for col in range(1, sheet.max_column + 1):
        header = _header_text(sheet, col, first_data_row)
        if "指数" not in header:
            continue
        if not any(token in header for token in year_tokens):
            continue
        candidates.append(col)
    assert candidates, f"IPSS index column for {year} was not found"
    # A strict semantic match should identify exactly one index column per year.
    assert len(candidates) == 1, f"IPSS index column for {year} is ambiguous: {candidates}"
    return candidates[0]


def _find_national_row(sheet, name_col: int) -> int:
    rows = []
    for row in range(1, sheet.max_row + 1):
        if _norm(sheet.cell(row, name_col).value) in {"全国", "全国計"}:
            rows.append(row)
    assert rows, "IPSS national row was not found"
    return rows[0]


def build_live_ipss_indices() -> tuple[dict[str, dict[str, float]], dict[str, float]]:
    wb = _fetch_workbook()
    sheet, name_col, pref_rows = _find_prefecture_block(wb)
    first_data_row = min(pref_rows.values())
    col_2035 = _find_index_column(sheet, first_data_row, 2035)
    col_2050 = _find_index_column(sheet, first_data_row, 2050)
    national_row = _find_national_row(sheet, name_col)

    nat_2035 = _number(sheet.cell(national_row, col_2035).value)
    nat_2050 = _number(sheet.cell(national_row, col_2050).value)
    assert nat_2035 == 77.8, f"IPSS national 2035 index must be 77.8, got {nat_2035}"
    assert nat_2050 == 69.2, f"IPSS national 2050 index must be 69.2, got {nat_2050}"

    live: dict[str, dict[str, float]] = {}
    for name, row in pref_rows.items():
        idx_2035 = _number(sheet.cell(row, col_2035).value)
        idx_2050 = _number(sheet.cell(row, col_2050).value)
        assert idx_2035 is not None and 0 < idx_2035 <= 150, f"{name}: invalid 2035 index {idx_2035}"
        assert idx_2050 is not None and 0 < idx_2050 <= 150, f"{name}: invalid 2050 index {idx_2050}"
        live[name] = {
            "child_population_index_2035": round(idx_2035, 1),
            "child_population_index_2050": round(idx_2050, 1),
        }

    assert len(live) == 47
    print(
        f"[IPSS Live Audit] sheet={sheet.title!r}, prefectures=47, "
        f"2035_col={col_2035}, 2050_col={col_2050}, national=77.8/69.2"
    )
    return live, {
        "child_population_index_2035": nat_2035,
        "child_population_index_2050": nat_2050,
    }


def test_ipss_future_child_population_ground_truth() -> None:
    master = json.loads(MASTER_PATH.read_text(encoding="utf-8"))
    assert master["source"]["xlsx_url"] == IPSS_XLSX_URL
    assert master["source"]["base_year"] == 2020
    assert len(master["prefectures"]) == 47

    committed_by_name = {item["prefecture_name"]: item for item in master["prefectures"]}
    assert set(committed_by_name) == PREFECTURE_SET

    live, national = build_live_ipss_indices()
    assert national == master["national"]

    for name in PREFECTURE_NAMES:
        committed = committed_by_name[name]
        assert live[name]["child_population_index_2035"] == committed["child_population_index_2035"], (
            f"{name}: 2035 live={live[name]['child_population_index_2035']} "
            f"committed={committed['child_population_index_2035']}"
        )
        assert live[name]["child_population_index_2050"] == committed["child_population_index_2050"], (
            f"{name}: 2050 live={live[name]['child_population_index_2050']} "
            f"committed={committed['child_population_index_2050']}"
        )


if __name__ == "__main__":
    test_ipss_future_child_population_ground_truth()
    print("ALL 47 PREFECTURES IPSS FUTURE CHILD POPULATION INDICES PASSED LIVE OFFICIAL GROUND TRUTH!")
